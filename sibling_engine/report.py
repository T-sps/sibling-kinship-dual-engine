"""Word 报告自动生成（非核心模块），依据国标附录 B 格式。

依赖 python-docx。模板细节待按附录 B 调整填充。
"""
from __future__ import annotations


def render_report(verdict: dict, output_path: str) -> None:
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("需安装 python-docx: pip install python-docx")

    doc = Document()
    doc.add_heading("全同胞关系鉴定意见书", level=1)

    # TODO: 按国标附录 B 填充：
    #   - 基因型比对表
    #   - 统计参数（CIBS / CFSI / Log10(CFSI)）
    #   - 判定结论
    #   - 可视化图表（多体热力图等）
    doc.add_paragraph(f"鉴定结论：{verdict.get('conclusion', '(待填充)')}")

    doc.save(output_path)
